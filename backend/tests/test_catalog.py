from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from PIL import Image
from rest_framework.test import APIClient

from apps.catalog.models import Category, IndustrySegment, Product, ProductImage
from apps.tenancy.models import Membership, Tenant, TenantDomain


@pytest.fixture
def catalog_context(db):
    tenant_a = Tenant.objects.create(
        name="Empresa A",
        slug="catalog-empresa-a",
        headline="Empresa A",
        logo_path="/static/demo/empresa-a.svg",
        featured_offerings=[],
    )
    tenant_b = Tenant.objects.create(
        name="Empresa B",
        slug="catalog-empresa-b",
        headline="Empresa B",
        logo_path="/static/demo/empresa-b.svg",
        featured_offerings=[],
    )
    TenantDomain.objects.create(tenant=tenant_a, hostname="catalog-a.localhost", is_primary=True, is_verified=True)
    TenantDomain.objects.create(tenant=tenant_b, hostname="catalog-b.localhost", is_primary=True, is_verified=True)
    owner = get_user_model().objects.create_user("catalog-owner", "owner@example.test", "password")
    viewer = get_user_model().objects.create_user("catalog-viewer", "viewer@example.test", "password")
    Membership.objects.create(tenant=tenant_a, user=owner, role=Membership.Role.OWNER)
    Membership.objects.create(tenant=tenant_a, user=viewer, role=Membership.Role.VIEWER)
    category_a = Category.objects.create(tenant=tenant_a, name="Abarrotes")
    category_b = Category.objects.create(tenant=tenant_b, name="Insumos")
    return tenant_a, tenant_b, owner, viewer, category_a, category_b


def authenticated_client(user):
    client = APIClient()
    client.force_authenticate(user)
    return client


def catalog_csv(*rows):
    header = "sku,nombre,categoria,descripcion,formato,precio,disponible,publicar\n"
    return SimpleUploadedFile("catalogo.csv", (header + "\n".join(rows) + "\n").encode(), content_type="text/csv")


def catalog_xlsx(*rows, headers=("sku", "nombre", "categoria", "descripcion", "formato", "precio", "disponible", "publicar")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Productos"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return SimpleUploadedFile("catalogo.xlsx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def product_png():
    content = BytesIO()
    Image.new("RGB", (80, 60), "#a88b69").save(content, format="PNG")
    return SimpleUploadedFile("cafe.png", content.getvalue(), content_type="image/png")


@pytest.mark.django_db
def test_product_image_upload_and_publication_are_tenant_scoped(catalog_context):
    tenant_a, _, owner, viewer, category_a, _ = catalog_context
    product = Product.objects.create(tenant=tenant_a, category=category_a, name="Café", sku="CAFE-1")
    upload_url = reverse("catalog-product-image", args=[product.id])
    public_url = reverse("public-product-image", args=[product.id])
    owner_client = authenticated_client(owner)

    assert authenticated_client(viewer).post(upload_url, {"image": product_png()}, format="multipart", HTTP_HOST="catalog-a.localhost").status_code == 403
    assert owner_client.post(upload_url, {"image": product_png()}, format="multipart", HTTP_HOST="catalog-b.localhost").status_code == 403
    assert APIClient().get(public_url, HTTP_HOST="catalog-a.localhost").status_code == 404

    uploaded = owner_client.post(upload_url, {"image": product_png()}, format="multipart", HTTP_HOST="catalog-a.localhost")
    assert uploaded.status_code == 200
    assert uploaded.json()["image_path"].startswith(upload_url)
    assert ProductImage.objects.filter(product=product).exists()
    assert owner_client.get(upload_url, HTTP_HOST="catalog-a.localhost")["Content-Type"] == "image/webp"
    assert APIClient().get(public_url, HTTP_HOST="catalog-a.localhost").status_code == 404

    product.is_published = True
    product.save(update_fields=["is_published"])
    landing = APIClient().get(reverse("public-landing"), HTTP_HOST="catalog-a.localhost")
    assert landing.json()["products"][0]["image_path"].startswith(public_url)
    public = APIClient().get(public_url, HTTP_HOST="catalog-a.localhost")
    assert public.status_code == 200
    assert public["Content-Type"] == "image/webp"
    assert APIClient().get(public_url, HTTP_HOST="catalog-b.localhost").status_code == 404

    product.is_available = False
    product.save(update_fields=["is_available"])
    assert APIClient().get(public_url, HTTP_HOST="catalog-a.localhost").status_code == 404
    product.is_available = True
    product.save(update_fields=["is_available"])
    category_a.is_active = False
    category_a.save(update_fields=["is_active"])
    assert APIClient().get(public_url, HTTP_HOST="catalog-a.localhost").status_code == 404
    category_a.is_active = True
    category_a.save(update_fields=["is_active"])

    removed = owner_client.delete(upload_url, HTTP_HOST="catalog-a.localhost")
    assert removed.status_code == 200
    assert removed.json()["image_path"] is None
    assert APIClient().get(public_url, HTTP_HOST="catalog-a.localhost").status_code == 404


@pytest.mark.django_db
def test_product_image_rejects_invalid_upload(catalog_context):
    tenant_a, _, owner, _, category_a, _ = catalog_context
    product = Product.objects.create(tenant=tenant_a, category=category_a, name="Café", sku="CAFE-1")
    response = authenticated_client(owner).post(
        reverse("catalog-product-image", args=[product.id]),
        {"image": SimpleUploadedFile("fake.png", b"not-an-image", content_type="image/png")},
        format="multipart",
        HTTP_HOST="catalog-a.localhost",
    )
    assert response.status_code == 400
    assert not ProductImage.objects.filter(product=product).exists()


@pytest.mark.django_db
def test_owner_creates_product_only_inside_resolved_tenant(catalog_context):
    tenant_a, _, owner, _, category_a, _ = catalog_context
    response = authenticated_client(owner).post(
        reverse("catalog-product-list"),
        {
            "category": str(category_a.id),
            "name": "Café premium",
            "sku": "CAFE-001",
            "format": "Caja 12 unidades",
            "price": "24990.00",
            "is_available": True,
            "is_published": False,
        },
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )

    assert response.status_code == 201
    product = Product.objects.get()
    assert product.tenant == tenant_a
    assert product.price == Decimal("24990.00")
    assert response.json()["category_name"] == "Abarrotes"


@pytest.mark.django_db
def test_product_target_industries_use_canonical_codes(catalog_context):
    _, _, owner, _, category, _ = catalog_context
    api = authenticated_client(owner)
    valid = api.post(
        reverse("catalog-product-list"),
        {
            "category": str(category.id), "name": "Salmón", "sku": "SAL-1",
            "target_industries": [IndustrySegment.FOOD_SERVICE, IndustrySegment.HOSPITALITY],
        },
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )
    assert valid.status_code == 201
    assert valid.json()["target_industry_labels"] == [
        "Restaurantes, restobares y cafeterías", "Hoteles y alojamientos",
    ]

    invalid = api.patch(
        reverse("catalog-product-detail", args=[valid.json()["id"]]),
        {"target_industries": ["Restaurant", "Restobar"]},
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )
    assert invalid.status_code == 400


@pytest.mark.django_db
def test_industry_taxonomy_is_served_from_backend(catalog_context):
    _, _, owner, _, _, _ = catalog_context
    response = authenticated_client(owner).get(
        reverse("catalog-industry-segments"), HTTP_HOST="catalog-a.localhost"
    )
    assert response.status_code == 200
    assert response.json()[0] == {
        "value": IndustrySegment.FOOD_SERVICE,
        "label": "Restaurantes, restobares y cafeterías",
    }


@pytest.mark.django_db
def test_catalog_rejects_category_from_another_tenant(catalog_context):
    _, _, owner, _, _, category_b = catalog_context
    response = authenticated_client(owner).post(
        reverse("catalog-product-list"),
        {"category": str(category_b.id), "name": "Producto cruzado", "sku": "CROSS-1"},
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )

    assert response.status_code == 400
    assert Product.objects.count() == 0


@pytest.mark.django_db
def test_catalog_rejects_negative_price(catalog_context):
    _, _, owner, _, category_a, _ = catalog_context
    response = authenticated_client(owner).post(
        reverse("catalog-product-list"),
        {"category": str(category_a.id), "name": "Producto", "sku": "NEG-1", "price": "-1.00"},
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )

    assert response.status_code == 400
    assert "price" in response.json()


@pytest.mark.django_db
def test_catalog_lists_only_current_tenant_products(catalog_context):
    tenant_a, tenant_b, owner, _, category_a, category_b = catalog_context
    Product.objects.create(tenant=tenant_a, category=category_a, name="Producto A", sku="A-1")
    Product.objects.create(tenant=tenant_b, category=category_b, name="Producto B", sku="B-1")

    response = authenticated_client(owner).get(reverse("catalog-product-list"), HTTP_HOST="catalog-a.localhost")

    assert response.status_code == 200
    assert [product["name"] for product in response.json()] == ["Producto A"]


@pytest.mark.django_db
def test_viewer_can_read_but_cannot_modify_catalog(catalog_context):
    _, _, _, viewer, category_a, _ = catalog_context
    product = Product.objects.create(tenant=category_a.tenant, category=category_a, name="Producto", sku="P-1")
    client = authenticated_client(viewer)

    assert client.get(reverse("catalog-product-list"), HTTP_HOST="catalog-a.localhost").status_code == 200
    response = client.patch(
        reverse("catalog-product-detail", args=[product.id]),
        {"is_published": True},
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )
    assert response.status_code == 403
    product.refresh_from_db()
    assert product.is_published is False


@pytest.mark.django_db
def test_public_landing_contains_only_published_available_products(catalog_context):
    tenant_a, _, _, _, category_a, _ = catalog_context
    Product.objects.create(
        tenant=tenant_a,
        category=category_a,
        name="Producto publicado",
        sku="PUB-1",
        price="15990.00",
        is_published=True,
    )
    Product.objects.create(tenant=tenant_a, category=category_a, name="Producto borrador", sku="DRAFT-1")

    response = APIClient().get(reverse("public-landing"), HTTP_HOST="catalog-a.localhost")

    assert response.status_code == 200
    assert [product["name"] for product in response.json()["products"]] == ["Producto publicado"]
    assert response.json()["products"][0]["category"] == "Abarrotes"


@pytest.mark.django_db
def test_csv_preview_then_confirmation_is_transactional(catalog_context):
    tenant_a, _, owner, _, _, _ = catalog_context
    client = authenticated_client(owner)
    preview = client.post(
        reverse("catalog-import-preview"),
        {"file": catalog_csv("NUEVO-1,Producto nuevo,Nueva categoría,Descripción,Caja 10,12500,sí,no")},
        format="multipart",
        HTTP_HOST="catalog-a.localhost",
    )

    assert preview.status_code == 200
    assert preview.json()["valid"] is True
    assert preview.json()["summary"] == {"total": 1, "create": 1, "update": 0, "errors": 0, "warnings": 0}
    assert Product.objects.count() == 0

    confirmation = client.post(
        reverse("catalog-import-confirm"),
        {"token": preview.json()["token"]},
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )
    assert confirmation.status_code == 200
    assert confirmation.json() == {"created": 1, "updated": 0, "total": 1}
    product = Product.objects.get(tenant=tenant_a, sku="NUEVO-1")
    assert product.category.name == "Nueva categoría"
    assert product.price == Decimal("12500.00")


@pytest.mark.django_db
def test_csv_reimport_updates_by_sku_without_duplicates(catalog_context):
    tenant_a, _, owner, _, category_a, _ = catalog_context
    Product.objects.create(tenant=tenant_a, category=category_a, name="Nombre anterior", sku="SKU-1", price="100.00")
    client = authenticated_client(owner)
    preview = client.post(
        reverse("catalog-import-preview"),
        {"file": catalog_csv("SKU-1,Nombre actualizado,Abarrotes,,Caja,200,si,si")},
        format="multipart",
        HTTP_HOST="catalog-a.localhost",
    )
    assert preview.json()["rows"][0]["action"] == "actualizar"

    response = client.post(
        reverse("catalog-import-confirm"),
        {"token": preview.json()["token"]},
        format="json",
        HTTP_HOST="catalog-a.localhost",
    )

    assert response.json() == {"created": 0, "updated": 1, "total": 1}
    assert Product.objects.filter(tenant=tenant_a, sku="SKU-1").count() == 1
    product = Product.objects.get(tenant=tenant_a, sku="SKU-1")
    assert product.name == "Nombre actualizado"
    assert product.is_published is True


@pytest.mark.django_db
def test_invalid_csv_returns_row_errors_and_cannot_be_confirmed(catalog_context):
    _, _, owner, _, _, _ = catalog_context
    client = authenticated_client(owner)
    response = client.post(
        reverse("catalog-import-preview"),
        {"file": catalog_csv("BAD-1,,Abarrotes,,,-20,quizás,sí")},
        format="multipart",
        HTTP_HOST="catalog-a.localhost",
    )

    assert response.status_code == 200
    assert response.json()["valid"] is False
    assert response.json()["token"] is None
    assert response.json()["summary"]["errors"] == 1
    assert len(response.json()["rows"][0]["errors"]) == 3
    assert Product.objects.count() == 0


@pytest.mark.django_db
def test_import_preview_token_cannot_cross_tenants(catalog_context):
    _, tenant_b, owner, _, _, _ = catalog_context
    user_b = get_user_model().objects.create_user("owner-b", "owner-b@example.test", "password")
    Membership.objects.create(tenant=tenant_b, user=user_b, role=Membership.Role.OWNER)
    client = authenticated_client(owner)
    preview = client.post(
        reverse("catalog-import-preview"),
        {"file": catalog_csv("SAFE-1,Producto,Abarrotes,,Caja,100,sí,no")},
        format="multipart",
        HTTP_HOST="catalog-a.localhost",
    )
    client.force_authenticate(user_b)

    response = client.post(
        reverse("catalog-import-confirm"),
        {"token": preview.json()["token"]},
        format="json",
        HTTP_HOST="catalog-b.localhost",
    )
    assert response.status_code == 400
    assert Product.objects.count() == 0


@pytest.mark.django_db
def test_catalog_xlsx_template_has_blank_products_sheet_and_instructions(catalog_context):
    _, _, owner, _, _, _ = catalog_context
    response = authenticated_client(owner).get(reverse("catalog-import-template"), HTTP_HOST="catalog-a.localhost")

    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="plantilla-catalogo.xlsx"'
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Productos", "Cómo completar"]
    assert [cell.value for cell in workbook["Productos"][1]] == ["sku", "nombre", "categoria", "descripcion", "formato", "precio", "disponible", "publicar"]
    assert workbook["Productos"]["A2"].value is None
    assert workbook["Productos"]["G2"].value is None


@pytest.mark.django_db
def test_xlsx_preview_skips_blank_rows_and_imports_pesos(catalog_context):
    tenant_a, _, owner, _, _, _ = catalog_context
    client = authenticated_client(owner)
    upload = catalog_xlsx(
        (None, None, None, None, None, None, None, None),
        ("MAR-001", "Camarón premium", "Mariscos", "Congelado", "Caja 10 kg", 24990, "sí", "no"),
    )
    preview = client.post(reverse("catalog-import-preview"), {"file": upload}, format="multipart", HTTP_HOST="catalog-a.localhost")
    assert preview.status_code == 200
    assert preview.json()["valid"] is True
    assert preview.json()["summary"]["total"] == 1
    assert preview.json()["rows"][0]["row"] == 3
    assert preview.json()["rows"][0]["price"] == "24990.00"
    assert Product.objects.count() == 0
    confirmed = client.post(reverse("catalog-import-confirm"), {"token": preview.json()["token"]}, format="json", HTTP_HOST="catalog-a.localhost")
    assert confirmed.status_code == 200
    assert Product.objects.get(tenant=tenant_a, sku="MAR-001").price == Decimal("24990.00")


@pytest.mark.django_db
def test_xlsx_rejects_formulas_and_wrong_columns(catalog_context):
    _, _, owner, _, _, _ = catalog_context
    client = authenticated_client(owner)
    formula = catalog_xlsx(("MAR-002", "=1+1", "Mariscos", "", "Caja", 15000, "sí", "no"))
    preview = client.post(reverse("catalog-import-preview"), {"file": formula}, format="multipart", HTTP_HOST="catalog-a.localhost")
    assert preview.status_code == 200
    assert preview.json()["valid"] is False
    assert preview.json()["token"] is None
    assert "fórmulas" in preview.json()["rows"][0]["errors"][0]
    wrong_headers = catalog_xlsx(("MAR-003", "Atún", "Pescados", "", "Caja", 15000, "sí", "no"), headers=("codigo", "nombre", "categoria", "descripcion", "formato", "precio", "disponible", "publicar"))
    rejected = client.post(reverse("catalog-import-preview"), {"file": wrong_headers}, format="multipart", HTTP_HOST="catalog-a.localhost")
    assert rejected.status_code == 400
    assert Product.objects.count() == 0


@pytest.mark.django_db
def test_downloaded_xlsx_template_can_be_filled_and_previewed(catalog_context):
    _, _, owner, _, _, _ = catalog_context
    client = authenticated_client(owner)
    template = client.get(reverse("catalog-import-template"), HTTP_HOST="catalog-a.localhost")
    workbook = load_workbook(BytesIO(template.content))
    sheet = workbook["Productos"]
    for column, value in enumerate(("ATU-001", "Atún en trozos", "Pescados", "Congelado", "Caja 20", 15990, "sí", "no"), start=1):
        sheet.cell(2, column).value = value
    output = BytesIO()
    workbook.save(output)
    upload = SimpleUploadedFile("plantilla-catalogo.xlsx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    preview = client.post(reverse("catalog-import-preview"), {"file": upload}, format="multipart", HTTP_HOST="catalog-a.localhost")
    assert preview.status_code == 200
    assert preview.json()["valid"] is True
    assert preview.json()["summary"]["total"] == 1


@pytest.mark.django_db
def test_corrupt_xlsx_has_clear_error(catalog_context):
    _, _, owner, _, _, _ = catalog_context
    upload = SimpleUploadedFile("roto.xlsx", b"no es excel", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response = authenticated_client(owner).post(reverse("catalog-import-preview"), {"file": upload}, format="multipart", HTTP_HOST="catalog-a.localhost")
    assert response.status_code == 400
    assert "Excel no es válido" in response.json()["detail"]
