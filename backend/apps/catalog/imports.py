import csv
import io
import secrets
import time
import zipfile
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Category, Product

CSV_HEADERS = ("sku", "nombre", "categoria", "descripcion", "formato", "precio", "disponible", "publicar")
MAX_CSV_BYTES = 1_000_000
MAX_XLSX_BYTES = 2_000_000
MAX_XLSX_UNCOMPRESSED_BYTES = 12_000_000
MAX_CSV_ROWS = 500
PREVIEW_TTL_SECONDS = 30 * 60
TRUE_VALUES = {"1", "si", "sí", "true", "verdadero", "yes"}
FALSE_VALUES = {"0", "no", "false", "falso"}


class CatalogImportError(ValueError):
    pass


def csv_template() -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(CSV_HEADERS)
    writer.writerow(
        ("CAFE-001", "Café premium", "Abarrotes", "Café para negocios", "Caja 12 unidades", "24990", "sí", "no")
    )
    return output.getvalue()


def xlsx_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Productos"
    sheet.append(CSV_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H501"
    widths = (19, 34, 24, 48, 27, 17, 16, 16)
    for index, width in enumerate(widths, start=1):
        column = sheet.cell(1, index)
        column.fill = PatternFill("solid", fgColor="0D766E")
        column.font = Font(bold=True, color="FFFFFF")
        column.alignment = Alignment(vertical="center")
        sheet.column_dimensions[column.column_letter].width = width
    sheet.row_dimensions[1].height = 26
    for row in range(2, 502):
        sheet.cell(row, 1).number_format = "@"
        sheet.cell(row, 6).number_format = '#,##0'
    yes_no = DataValidation(type="list", formula1='"sí,no"', allow_blank=True)
    yes_no.error = "Selecciona sí o no."
    yes_no.errorTitle = "Valor no válido"
    yes_no.showErrorMessage = True
    sheet.add_data_validation(yes_no)
    yes_no.add("G2:H501")

    guide = workbook.create_sheet("Cómo completar")
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 72
    instructions = [
        ("Plantilla de productos", "Completa una fila por producto en la hoja Productos. La primera fila no se modifica."),
        ("sku", "Código único por producto. Si ya existe en tu empresa, se actualizará."),
        ("nombre", "Nombre del producto para el catálogo."),
        ("categoria", "Nombre de la categoría. Se creará si todavía no existe."),
        ("descripcion y formato", "Datos opcionales. Ejemplo de formato: Caja 10 unidades."),
        ("precio", "Precio en pesos chilenos, sin signo $ ni separador de miles. Ejemplo: 24990. Déjalo vacío para cotizar."),
        ("disponible", "Elige sí o no."),
        ("publicar", "Elige sí para mostrarlo en la landing o no para mantenerlo como borrador."),
        ("Imágenes", "Después de importar, sube las fotos de cada producto en la sección Marketing."),
        ("Ejemplo", "CAFE-001 | Café premium | Abarrotes | Café para negocios | Caja 12 unidades | 24990 | sí | no"),
        ("Antes de importar", "Guarda este archivo como .xlsx. En la app revisa la vista previa antes de confirmar."),
    ]
    for row_number, (title, detail) in enumerate(instructions, start=1):
        guide.cell(row_number, 1, title).font = Font(bold=True, color="0D766E")
        guide.cell(row_number, 2, detail).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row_number].height = 34 if row_number in {1, 6, 10} else 28
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _text(value) -> str:
    return str(value if value is not None else "").strip()


def _boolean(value: str, field: str, errors: list[str]) -> bool | None:
    normalized = value.casefold()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    errors.append(f"{field} debe ser sí o no.")
    return None


def _price(value: str, errors: list[str]) -> str | None:
    if not value:
        return None
    try:
        parsed = Decimal(value.replace(".", "").replace(",", ".") if "," in value else value)
    except InvalidOperation:
        errors.append("precio debe ser un número válido.")
        return None
    if parsed < 0:
        errors.append("precio no puede ser negativo.")
        return None
    if parsed.as_tuple().exponent < -2 or parsed >= Decimal("10000000000"):
        errors.append("precio debe tener máximo dos decimales y diez dígitos enteros.")
        return None
    return f"{parsed:.2f}"


def preview_csv(uploaded_file, tenant) -> dict:
    if uploaded_file.size > MAX_CSV_BYTES:
        raise CatalogImportError("El archivo supera el máximo permitido de 1 MB.")
    if not uploaded_file.name.lower().endswith(".csv"):
        raise CatalogImportError("Selecciona un archivo con extensión .csv.")
    try:
        content = uploaded_file.read().decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CatalogImportError("El archivo debe usar codificación UTF-8.") from exc

    reader = csv.DictReader(io.StringIO(content, newline=""))
    headers = tuple(reader.fieldnames or ())
    if headers != CSV_HEADERS:
        raise CatalogImportError(f"Las columnas deben ser exactamente: {', '.join(CSV_HEADERS)}.")

    return _preview_rows(((number, source, False) for number, source in enumerate(reader, start=2)), tenant)


def preview_xlsx(uploaded_file, tenant) -> dict:
    if uploaded_file.size > MAX_XLSX_BYTES:
        raise CatalogImportError("El archivo Excel supera el máximo permitido de 2 MB.")
    if not uploaded_file.name.lower().endswith(".xlsx"):
        raise CatalogImportError("Selecciona un archivo .xlsx.")
    content = uploaded_file.read()
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 100 or sum(entry.file_size for entry in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise CatalogImportError("El archivo Excel es demasiado grande para procesarlo.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except CatalogImportError:
        raise
    except Exception as exc:
        raise CatalogImportError("El archivo Excel no es válido. Usa una plantilla .xlsx sin macros.") from exc
    try:
        if "Productos" not in workbook.sheetnames:
            raise CatalogImportError("El archivo debe incluir una hoja llamada Productos.")
        sheet = workbook["Productos"]
        if sheet.max_row > MAX_CSV_ROWS + 1:
            raise CatalogImportError(f"La hoja Productos puede contener como máximo {MAX_CSV_ROWS} filas de productos.")
        headers = tuple(_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=8)))
        if headers != CSV_HEADERS or sheet.max_column != len(CSV_HEADERS):
            raise CatalogImportError(f"Las columnas de Productos deben ser exactamente: {', '.join(CSV_HEADERS)}.")

        def source_rows():
            for cells in sheet.iter_rows(min_row=2, max_col=len(CSV_HEADERS)):
                source = dict(zip(CSV_HEADERS, (cell.value for cell in cells)))
                if all(not _text(value) for value in source.values()):
                    continue
                yield cells[0].row, source, any(cell.data_type == "f" for cell in cells)

        return _preview_rows(source_rows(), tenant)
    except CatalogImportError:
        raise
    except Exception as exc:
        raise CatalogImportError("No fue posible leer la hoja Productos. Revisa que el archivo .xlsx no esté dañado.") from exc
    finally:
        workbook.close()


def preview_catalog(uploaded_file, tenant) -> dict:
    if uploaded_file.name.lower().endswith(".xlsx"):
        return preview_xlsx(uploaded_file, tenant)
    if uploaded_file.name.lower().endswith(".csv"):
        return preview_csv(uploaded_file, tenant)
    raise CatalogImportError("Selecciona un archivo .xlsx o .csv.")


def _preview_rows(sources, tenant) -> dict:
    existing_skus = set(Product.objects.filter(tenant=tenant).values_list("sku", flat=True))
    rows = []
    normalized_rows = []
    seen_skus = set()
    for row_number, source, has_formula in sources:
        if len(rows) >= MAX_CSV_ROWS:
            raise CatalogImportError(f"El archivo puede contener como máximo {MAX_CSV_ROWS} productos.")
        errors: list[str] = ["Las fórmulas no están permitidas. Ingresa valores directamente."] if has_formula else []
        warnings: list[str] = []
        sku = _text(source["sku"])
        name = _text(source["nombre"])
        category = _text(source["categoria"])
        description = _text(source["descripcion"])
        product_format = _text(source["formato"])
        if not sku:
            errors.append("sku es obligatorio.")
        elif len(sku) > 80:
            errors.append("sku supera 80 caracteres.")
        elif sku.casefold() in seen_skus:
            errors.append("sku está repetido en el archivo.")
        seen_skus.add(sku.casefold())
        if not name:
            errors.append("nombre es obligatorio.")
        elif len(name) > 160:
            errors.append("nombre supera 160 caracteres.")
        if not category:
            errors.append("categoria es obligatoria.")
        elif len(category) > 120:
            errors.append("categoria supera 120 caracteres.")
        if len(description) > 1000:
            errors.append("descripcion supera 1000 caracteres.")
        if len(product_format) > 120:
            errors.append("formato supera 120 caracteres.")
        if not product_format:
            warnings.append("formato está vacío.")
        price = _price(_text(source["precio"]), errors)
        available = _boolean(_text(source["disponible"]), "disponible", errors)
        published = _boolean(_text(source["publicar"]), "publicar", errors)
        action = "actualizar" if sku in existing_skus else "crear"
        normalized = {
            "sku": sku,
            "name": name,
            "category": category,
            "description": description,
            "format": product_format,
            "price": price,
            "is_available": available,
            "is_published": published,
        }
        rows.append(
            {
                "row": row_number,
                "sku": sku,
                "name": name,
                "category": category,
                "price": price,
                "action": action,
                "errors": errors,
                "warnings": warnings,
            }
        )
        normalized_rows.append(normalized)

    if not rows:
        raise CatalogImportError("El archivo no contiene productos.")
    error_count = sum(bool(row["errors"]) for row in rows)
    return {
        "valid": error_count == 0,
        "rows": rows,
        "normalized_rows": normalized_rows,
        "summary": {
            "total": len(rows),
            "create": sum(row["action"] == "crear" for row in rows),
            "update": sum(row["action"] == "actualizar" for row in rows),
            "errors": error_count,
            "warnings": sum(bool(row["warnings"]) for row in rows),
        },
    }


def save_preview(session, tenant_id, normalized_rows) -> str:
    token = secrets.token_urlsafe(24)
    previews = session.get("catalog_import_previews", {})
    previews[token] = {"tenant_id": str(tenant_id), "created_at": int(time.time()), "rows": normalized_rows}
    session["catalog_import_previews"] = previews
    session.modified = True
    return token


def confirm_preview(session, tenant, token: str) -> dict:
    previews = session.get("catalog_import_previews", {})
    preview = previews.get(token)
    if not preview or preview["tenant_id"] != str(tenant.id):
        raise CatalogImportError("La vista previa no existe o pertenece a otra empresa.")
    if int(time.time()) - preview["created_at"] > PREVIEW_TTL_SECONDS:
        previews.pop(token, None)
        session["catalog_import_previews"] = previews
        raise CatalogImportError("La vista previa expiró. Vuelve a cargar el archivo.")

    created = updated = 0
    with transaction.atomic():
        for row in preview["rows"]:
            category, _ = Category.objects.get_or_create(tenant=tenant, name=row["category"])
            _, was_created = Product.objects.update_or_create(
                tenant=tenant,
                sku=row["sku"],
                defaults={
                    "category": category,
                    "name": row["name"],
                    "description": row["description"],
                    "format": row["format"],
                    "price": row["price"],
                    "is_available": row["is_available"],
                    "is_published": row["is_published"],
                },
            )
            created += was_created
            updated += not was_created
    previews.pop(token, None)
    session["catalog_import_previews"] = previews
    session.modified = True
    return {"created": created, "updated": updated, "total": created + updated}
